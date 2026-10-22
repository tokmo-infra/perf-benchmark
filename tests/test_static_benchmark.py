"""Native static benchmarking: timing, unknown usage, invalid config and persisted facts."""
import asyncio
import json
from pathlib import Path
from zipfile import ZipFile
import pytest
from aiohttp import web
from tokenizers import Tokenizer
from tokenizers.models import WordLevel
from tokenizers.pre_tokenizers import Whitespace
from transformers import PreTrainedTokenizerFast
from perf_benchmark.static_benchmark.config import load_static_benchmark_config
from perf_benchmark.static_benchmark.runner import StaticBenchmarkRunner
from perf_benchmark.static_benchmark.token_sampler import VocabularyTokenSampler
from perf_benchmark.static_benchmark.cli import run_static_cli
from perf_benchmark.plugins.saver.sqlite import SQLiteSaver


@pytest.fixture(autouse=True)
def private_sealing_root(tmp_path, monkeypatch):
    from perf_benchmark.plugins.saver.sqlite import archive
    root = tmp_path / 'sealing-coordination'
    root.mkdir()
    monkeypatch.setattr(archive, '_sealing_root', lambda: root)


@pytest.fixture
def static_tokenizer(tmp_path):
    words = ["[UNK]", "user", "assistant"] + [f"word{i}" for i in range(32)]
    backend = Tokenizer(WordLevel({word: i for i, word in enumerate(words)}, unk_token="[UNK]"))
    backend.pre_tokenizer = Whitespace()
    tokenizer = PreTrainedTokenizerFast(tokenizer_object=backend, unk_token="[UNK]")
    tokenizer.chat_template = "{% for message in messages %}user {{ message['content'] }} {% endfor %}assistant"
    path = tmp_path / "tokenizer"
    tokenizer.save_pretrained(path)
    return path


def test_sampler_requires_real_local_tokenizer_and_counts_chat_template(static_tokenizer):
    with pytest.raises(ValueError, match="requires a local tokenizer"):
        VocabularyTokenSampler(None)
    sampler = VocabularyTokenSampler(static_tokenizer)
    first = sampler.sample_prompt(16, 42, adjust_for_chat_template=True)
    second = sampler.sample_prompt(16, 43, adjust_for_chat_template=True)
    assert len(sampler._tokenizer.encode_chat([{"role": "user", "content": first}])) == 16
    assert first == sampler.sample_prompt(16, 42, adjust_for_chat_template=True)
    assert first != second
    assert len(sampler.fingerprint) == 64


@pytest.mark.parametrize("values", [
    {"prefill": {"matrix": {16: [0]}}},
    {"prefill": {"matrix": {16: [1, 1]}}},
    {"prefill": {"enabled": "false"}},
    {"service": {"stream": False}},
    {"service": {"extra_body": {"max_tokens": 1}}},
    {"decode": {"concurrencies": [-1]}},
    {"decode": {"requests_per_concurrency": 0}},
    {"output": {"excel_path": "same", "json_path": "same"}},
    {"prefil": {}},
])
def test_config_rejects_silent_workload_changes(values):
    with pytest.raises(ValueError):
        load_static_benchmark_config({'case':'invalid configuration test',**values})


@pytest.fixture
async def static_server():
    captured = []
    mode = {"usage": True, "status": 200}
    async def handle(request):
        body = await request.json()
        captured.append(body)
        if mode["status"] != 200:
            return web.json_response({"error": {"message": "rejected workload"}}, status=mode["status"])
        response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
        await response.prepare(request)
        async def send(data):
            await response.write(("data: "+json.dumps(data)+"\n\n").encode())
        await send({"choices": [{"delta": {"role": "assistant"}}]})
        await asyncio.sleep(.04)
        await send({"choices": [{"delta": {"reasoning_content": "Planning"}}]})
        await asyncio.sleep(.015)
        await send({"choices": [{"delta": {"content": "Code"}}]})
        await asyncio.sleep(.04)
        await send({"choices": [{"delta": {}, "finish_reason": "stop"}]})
        if mode["usage"]:
            await send({"choices": [], "usage": {"prompt_tokens": 16, "completion_tokens": 64,
                                                 "prompt_tokens_details": {"cached_tokens": 0}}})
        await response.write(b"data: [DONE]\n\n")
        return response
    app = web.Application()
    app.router.add_post("/v1/chat/completions", handle)
    async def info(request):
        return web.json_response({"dp_size": 1, "internal_states": [{"memory_usage": {"token_capacity": 123456}}]})
    async def metrics(request):
        return web.Response(text='sglang:hicache_host_total_tokens{model_name="mock",tp_rank="0"} 370368\n')
    app.router.add_get("/get_server_info", info)
    app.router.add_get("/metrics", metrics)
    runner = web.AppRunner(app)
    await runner.setup()
    site = web.TCPSite(runner, "127.0.0.1", 0)
    await site.start()
    try:
        yield f"http://127.0.0.1:{site._server.sockets[0].getsockname()[1]}/v1", captured, mode
    finally:
        await runner.cleanup()


async def test_native_static_matrix_persists_facts_and_excludes_role_and_usage_timing(
    tmp_path, static_tokenizer, static_server, monkeypatch,
):
    monkeypatch.setenv("PERF_BENCHMARK_PREPARATION_WORKERS", "1")
    url, captured, _ = static_server
    config = load_static_benchmark_config({
        'case':'Prefill与Decode端到端验证',
        "service": {"base_url": url, "model": "mock", "tokenizer_path": str(static_tokenizer)},
        "prefill": {"matrix": {16: [1, 2]}},
        "decode": {"concurrencies": [1], "rounds": 4, "requests_per_concurrency": 2,
                   "prompts": ["Write a small program"]},
        "output": {"excel_path": str(tmp_path / "result.xlsx")},
    })
    runner = StaticBenchmarkRunner(config)
    prefill, decode = await runner.run_all()
    assert len(prefill) == 2 and len(decode) == 1
    assert len(captured) == 17
    assert captured[0]["messages"] != captured[1]["messages"]
    assert all("max_completion_tokens" not in row for row in captured)
    assert all(row['temperature']==0.1 for row in captured)
    assert captured[-1]["reasoning_effort"] == "high"
    for result in (*prefill, *decode):
        assert result.status == "succeeded"
        assert result.ttft_mean_ms >= 30
        assert result.decode_speed_mean > 1000
        reader = SQLiteSaver.open_reader(result.database)
        assert reader.summary(result.run_id)["turn_count"] == result.rounds * result.requests_per_round
        assert len(reader.trials(result.run_id)) == 1
        assert result.eligible_samples("ttft") == result.total_requests
        assert "P99仅供观察" in result.tail_sample_note
        assert result.service_capacity["status"] == "observed"
        assert [p["capacity_tokens"] for p in result.service_capacity["pools"]] == [123456, 370368]
        assert Path(result.database + ".archive").is_dir()
    assert not list(runner.result_dir.glob(".static-input-*"))
    with ZipFile(config.output.excel_path) as workbook:
        from xml.etree import ElementTree as ET
        assert workbook.testzip() is None
        assert "指标可用性" in workbook.read("xl/workbook.xml").decode()
        assert "KV容量" in workbook.read("xl/workbook.xml").decode()
        assert "370368" in workbook.read("xl/worksheets/sheet5.xml").decode()
        ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        detail = ET.fromstring(workbook.read('xl/worksheets/sheet4.xml'))
        header = [cell.findtext('x:is/x:t', namespaces=ns)
                  for cell in detail.find('x:sheetData/x:row', ns)]
        assert '可用性' not in header and '缺失原因' in header


def test_round_quota_and_minimum():
    config = load_static_benchmark_config({'case':'round quota',"prefill": {"rounds": 4, "requests_per_concurrency": 64}})
    assert config.prefill.rounds == 4 and config.decode.rounds == 3
    for rounds in (0, 1, 2, True, 3.5):
        with pytest.raises(ValueError):
            load_static_benchmark_config({'case':'invalid rounds',"decode": {"rounds": rounds}})


@pytest.mark.parametrize('static_first', [True, False])
async def test_static_and_multiturn_share_excel_in_either_order(
    tmp_path, static_tokenizer, static_server, monkeypatch, stage_context, static_first,
):
    from openpyxl import load_workbook
    from tests.exporter.test_console import Reader
    from perf_benchmark.plugins.exporter.tabular import TabularExporter, parse_settings

    monkeypatch.setenv('PERF_BENCHMARK_PREPARATION_WORKERS', '1')
    url, _, _ = static_server
    path = tmp_path / 'combined.xlsx'

    class Results(Reader):
        location = str(tmp_path / 'facts.db')

        def request_page(self, run_id, **kwargs):
            return {'items': (), 'total': 0}

        def measurements(self, run_id, *, scope, scope_id):
            return {'evidence': {'case': run_id}} if scope == 'run' else {}

    exporter = TabularExporter(parse_settings({
        'output': str(tmp_path / 'tables'), 'excel_path': str(path)}))
    if not static_first:
        await exporter.export('multi-first', Results(), stage_context('exporter'))
    result_dirs = []
    for index in range(2):
        config = load_static_benchmark_config({
            'case': f'static-{index}',
            'service': {'base_url': url, 'model': 'mock', 'tokenizer_path': str(static_tokenizer)},
            'prefill': {'matrix': {16: [1]}},
            'decode': {'concurrencies': [1]},
            'output': {'excel_path': str(path)},
        })
        runner = StaticBenchmarkRunner(config)
        await runner.run_all()
        runner._checkpoint()  # Repeated checkpoints keep the same rows.
        result_dirs.append(runner.result_dir)
        await exporter.export(f'multi-{index}', Results(), stage_context('exporter'))
    assert result_dirs[0] != result_dirs[1]
    assert all((folder / 'prefill-16-c1' / 'run.db').is_file() for folder in result_dirs)
    book = load_workbook(path)
    try:
        assert {'summary', 'Prefill性能测试', 'Decode性能测试', 'KV容量'} <= set(book.sheetnames)
        for name in ('Prefill性能测试', 'Decode性能测试'):
            assert book[name].max_row == 3
            assert [book[name].cell(i, 1).value.split(' / ')[0] for i in (2, 3)] == ['static-0', 'static-1']
            assert book[name].row_dimensions[2].height == book[name].row_dimensions[3].height == 48
        assert book['summary'].max_row == (3 if static_first else 4)
        assert sum(row[0] == 'case' for row in book['测试概览与配置'].values) == 2
        assert all(sheet.freeze_panes == 'B2' for sheet in book)
    finally:
        book.close()


async def test_missing_usage_remains_unknown_and_failure_cli_returns_nonzero(
    tmp_path, static_tokenizer, static_server, monkeypatch,
):
    monkeypatch.setenv("PERF_BENCHMARK_PREPARATION_WORKERS", "1")
    url, _, mode = static_server
    mode["usage"] = False
    values = {'case':'missing usage and HTTP failure',"service": {"base_url": url, "model": "mock", "tokenizer_path": str(static_tokenizer)},
              "prefill": {"enabled": False}, "decode": {"concurrencies": [1]},
              "output": {"excel_path": str(tmp_path / "unknown.xlsx")}}
    runner = StaticBenchmarkRunner(load_static_benchmark_config(values))
    _, results = await runner.run_all()
    assert results[0].avg_actual_completion_tokens is None
    assert results[0].decode_speed_mean is None
    assert results[0].native_metrics["decode_throughput"]["statistic_metadata"]["am"]["availability"] == "unavailable"
    mode["status"] = 400
    values["output"]["excel_path"] = str(tmp_path / "failed.xlsx")
    path = tmp_path / "config.json"
    path.write_text(json.dumps(values))
    from argparse import Namespace
    assert await run_static_cli(Namespace(config=path)) == 1
    assert (tmp_path / "failed.xlsx").is_file()


async def test_dry_run_avoids_tokenizer_loading_and_network(tmp_path):
    from argparse import Namespace
    path = tmp_path / "config.json"
    path.write_text(json.dumps({'case':'dry run',"service": {"tokenizer_path": "/missing/tokenizer"},
                                "output": {"excel_path": str(tmp_path / "dry.xlsx")}}))
    assert await run_static_cli(Namespace(config=path, dry_run=True)) == 0
    assert not (tmp_path / "dry.xlsx").exists()
