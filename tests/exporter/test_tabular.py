import csv
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from perf_benchmark.core.models import Run, Trial, Session, TurnMeasurements, RequestMeasurements
from perf_benchmark.pipeline.contracts.results import (
    ComputedResultBatch,
    TurnResultRecord,
    AttemptResultRecord,
)
from perf_benchmark.pipeline.results.saver import SaverRuntime
from perf_benchmark.plugins.metrics.standard import StandardMetrics
from perf_benchmark.plugins.saver.sqlite import SQLiteSaver
from perf_benchmark.plugins.exporter.tabular import TabularExporter, parse_settings
from perf_benchmark.app.report import export_tables
from tests.metrics.test_maxtpm_billing import PRICES, occupancy, session, turn, trial
from tests.result_records import frozen_record


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def xlsx_rows(z, number):
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(z.read(f"xl/worksheets/sheet{number}.xml"))
    rows = []
    for row in root.findall("x:sheetData/x:row", ns):
        rows.append(
            [
                c.findtext("x:is/x:t", default="", namespaces=ns)
                if c.get("t") == "inlineStr"
                else c.findtext("x:v", default="", namespaces=ns)
                for c in row
            ]
        )
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def without_header_units(rows):
    return [{key.split(' (')[0]: value for key, value in row.items()} for row in rows]


def without_unit_columns(rows):
    return [{key: value for key, value in row.items() if not key.endswith('_unit')} for row in rows]


def test_revenue_columns_never_disappear_when_all_values_are_unknown():
    from perf_benchmark.plugins.exporter.tabular.render import (
        compact_summary, values, value_columns, REVENUE_METRICS, SUMMARY_METRICS, REASONLESS_METRICS,
    )
    header = value_columns(SUMMARY_METRICS)
    source = tuple(values({name: {} for name in ()}, SUMMARY_METRICS))
    _, labels, factory = compact_summary(('summary', header, lambda: iter([source])))
    result = dict(zip(labels, next(factory())))
    for name in REVENUE_METRICS:
        assert name in result and result[name] is None
        if name in REASONLESS_METRICS:
            assert name + '_reason' not in result
        else:
            assert result[name + '_reason'] == 'metric_not_recorded'
    assert not any(name.startswith('filtered_') for name in labels)


def test_maxtpm_million_unit_scales_values_once_without_changing_native_rows():
    from perf_benchmark.plugins.exporter.tabular.render import scale_maxtpm, units_in_headers, PartialValue
    original = [(12_716_578.304347826, 'tok/min'), (0, 'tok/min'), (None, 'tok/min'),
                (PartialValue(1_000_000), 'tok/min')]
    spec = ('summary', ('maxtpm', 'maxtpm_unit'), lambda: iter(original))
    scaled = scale_maxtpm(spec)
    values = list(scaled[2]())
    assert values[:3] == [(12.716578304347826, '百万token/M'), (0, '百万token/M'), (None, '百万token/M')]
    assert isinstance(values[3][0], PartialValue)
    assert list(scale_maxtpm(scaled)[2]()) == values
    folded = units_in_headers(scaled)
    assert folded[1] == ('maxtpm (百万token/M)',)
    assert original[0][0] == 12_716_578.304347826
    assert scale_maxtpm(folded) is folded


def test_compact_tables_omit_retired_columns_and_keep_rows_aligned():
    from perf_benchmark.plugins.exporter.tabular.render import tables
    from tests.exporter.test_console import Reader, _point

    class Compact(Reader):
        def measurements(self, *args, **kwargs):
            return {"evidence": {"conditions": {"stream": {"value": True}}}}

        def metric_windows(self, *args, **kwargs):
            return ({"window": "maxtpm:minute:0", "start_ns": 0,
                     "end_ns": 60_000_000_000, "duration_s": 60},)

        def metric_points(self, *args, **kwargs):
            return {name: _point("test-unit", {"value": 7},
                                 availability="partial", reason="partial_usage")
                    for name in ("failed", "cancelled", "success_rate", "maxtpm_cache_hit_rate",
                                 "cached_input_revenue_per_minute", "revenue_per_minute",
                                 "input_revenue_per_minute")}

        def request_page(self, *args, **kwargs):
            return {"items": ({"id": "request", "status": "failed",
                                "usage": {"prompt_tokens": 10, "completion_tokens": 2,
                                          "cached_tokens": 3}},), "total": 1}

    retired = {"stream", "trial_id", "maxtpm_cache_hit_rate_reason",
               "cached_input_revenue_per_minute_reason", "revenue_per_minute_reason"}
    retired.update(name + suffix for name in ("failed", "cancelled", "success_rate")
                   for suffix in ("", "_unit", "_availability", "_reason"))
    for name, header, factory in tables("run", Compact()):
        assert not retired.intersection(header)
        rows = list(factory())
        assert rows and all(len(row) == len(header) for row in rows)
        first = dict(zip(header, rows[0]))
        assert first["run_id"] == "run"
        if name in {"summary", "minutes"}:
            assert first["revenue_per_minute"] == 7
            assert first["input_revenue_per_minute_reason"] == "partial_usage"
        else:
            assert first["request_id"] == "request"
            assert first["status"] == "failed"
            assert first["uncached_input_tokens"] == 7


def test_starred_income_stays_numeric_in_the_existing_four_columns(tmp_path):
    from perf_benchmark.plugins.exporter.tabular.render import (
        PartialValue, REVENUE_METRICS, write_xlsx, csv_value,
    )
    path = tmp_path / 'income.xlsx'
    row = tuple(PartialValue(v) for v in (3.0, 2.0, 0.0, 1.0))
    write_xlsx(path, [('summary', REVENUE_METRICS, lambda: iter([row]))])
    with ZipFile(path) as z:
        result = xlsx_rows(z, 1)[0]
        assert tuple(result) == REVENUE_METRICS
        assert [float(result[key]) for key in REVENUE_METRICS] == list(row)
        assert len(ET.fromstring(z.read('xl/worksheets/sheet1.xml')).findall(
            './/{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row')[1]) == 4
        assert b'General&amp;quot;' not in z.read('xl/styles.xml')
        assert b'General&quot;*&quot;' in z.read('xl/styles.xml')
    assert csv_value(row[2]) == '0.0*'  # Known zero is retained, not omitted.


def test_report_column_projection_retains_missing_reasons_and_numeric_partial_values():
    from perf_benchmark.plugins.exporter.tabular.render import without_availability, PartialValue
    partial = PartialValue(12.5)
    header = ('case', 'availability', 'cost_availability', '可用性', 'value', 'reason')
    rows = [('partial', 'partial', 'partial', 'partial', partial, 'missing_cache'),
            ('unknown', 'unavailable', 'unavailable', 'unavailable', None, 'missing_usage'),
            ('zero', 'available', 'available', 'available', 0, None)]
    _, headings, factory = without_availability(('metrics', header, lambda: iter(rows)))
    assert headings == ('case', 'value', 'reason')
    assert list(factory()) == [('partial', partial, 'missing_cache'), ('unknown', None, 'missing_usage'),
                               ('zero', 0, None)]
    assert isinstance(list(factory())[0][1], PartialValue)


def test_xlsx_reads_each_factory_once_and_keeps_late_unicode_width(tmp_path):
    from perf_benchmark.plugins.exporter.tabular.render import write_xlsx, text_width
    calls, yielded = [], []
    label = "末尾完整配置名" * 12

    def rows():
        calls.append(1)
        assert len(calls) == 1, "export must not replay SQL/JSON for widths"
        for i in range(1000):
            yielded.append(i)
            yield (label if i == 999 else "short", i)

    path = tmp_path / "single-pass.xlsx"
    write_xlsx(path, [("summary", ("case", "value"), rows)])
    assert len(calls) == 1 and yielded == list(range(1000))
    with ZipFile(path) as z:
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        assert len(sheet.findall("x:sheetData/x:row", ns)) == 1001
        assert float(sheet.find("x:cols/x:col", ns).get("width")) == min(255, text_width(label) + 3)


def test_excel_first_column_fits_unicode_and_freezes_both_axes(tmp_path):
    from perf_benchmark.plugins.exporter.tabular.render import write_xlsx,text_width
    label='GLM5.3 '+('多轮缓存吞吐对照'*5)
    path=tmp_path/'case.xlsx'
    write_xlsx(path,[('summary',['case','value'],lambda:iter([(label,1)]))])
    with ZipFile(path) as z:
        ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        pane=root.find('x:sheetViews/x:sheetView/x:pane',ns)
        assert pane.attrib==dict(xSplit='1',ySplit='1',topLeftCell='B2',activePane='bottomRight',state='frozen')
        first=root.find('x:cols/x:col',ns)
        assert float(first.get('width'))>=text_width(label)+3
        assert root.find('x:sheetData/x:row[@r="2"]/x:c',ns).get('s') is None


def test_append_matching_headers_and_preserve_original_on_mismatch(tmp_path):
    from perf_benchmark.plugins.exporter.tabular.render import write_xlsx,PartialValue
    from perf_benchmark.plugins.exporter.tabular.append import append_workbook
    from openpyxl import load_workbook
    old,new,joined=tmp_path/'old.xlsx',tmp_path/'new.xlsx',tmp_path/'joined.xlsx'
    write_xlsx(old,[('summary',['case','cost'],lambda:iter([('before',1)]))])
    write_xlsx(new,[('summary',['case','cost'],lambda:iter([('after',PartialValue(2))]))])
    before=old.read_bytes()
    append_workbook(old,new,joined)
    book=load_workbook(joined)
    try:
        assert list(book['summary'].values)==[('case','cost'),('before',1),('after',2)]
        assert book['summary']['B3'].number_format=='General"*"'
        assert book['summary'].freeze_panes=='B2'
    finally:
        book.close()
    assert old.read_bytes()==before
    write_xlsx(new,[('summary',['different','cost'],lambda:iter([('bad',9)]))])
    with pytest.raises(ValueError,match='header mismatch'):
        append_workbook(old,new,joined)
    assert old.read_bytes()==before


async def test_explicit_excel_path_appends_trials_and_rejects_changed_headers(tmp_path, stage_context):
    from tests.exporter.test_console import Reader
    from openpyxl import load_workbook

    class Results(Reader):
        location=str(tmp_path/'facts.db')

        def request_page(self,run_id,**kwargs):
            return {'items':(), 'total':0}

        def measurements(self,run_id,*,scope,scope_id):
            return {'evidence':{'case':run_id}} if scope=='run' else {}

    path=tmp_path/'comparison.xlsx'
    settings=parse_settings({'output':str(tmp_path/'tables'),'excel_path':str(path)})
    exporter=TabularExporter(settings)
    await exporter.export('first',Results(),stage_context('exporter'))
    await exporter.export('second',Results(),stage_context('exporter'))
    book=load_workbook(path)
    try:
        assert [row[0] for row in list(book['summary'].values)[1:]]==['first','second']
        assert book['summary'].freeze_panes=='B2'
        book['summary']['A1']='wrong header'
        book.save(path)
    finally:
        book.close()
    before=path.read_bytes()
    with pytest.raises(ValueError,match='header mismatch'):
        await exporter.export('third',Results(),stage_context('exporter'))
    assert path.read_bytes()==before
    assert not list((tmp_path/'tables').glob('.tabular-*'))


def test_shared_workbook_adds_sheets_preserves_layout_and_publishes_atomically(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from openpyxl import load_workbook
    from perf_benchmark.plugins.exporter.tabular.render import write_xlsx
    from perf_benchmark.plugins.exporter.tabular.append import publish_workbook

    path = tmp_path / 'shared.xlsx'
    sources = [tmp_path / f'{i}.xlsx' for i in range(3)]
    for i, source in enumerate(sources):
        write_xlsx(source, [('summary', ['case', 'value'], lambda i=i: iter([(f'run-{i}', i)])),
                            (f'phase-{i}', ['case', 'value'], lambda: iter([('phase', 1)]))])
    with ThreadPoolExecutor(max_workers=3) as pool:
        list(pool.map(lambda source: publish_workbook(source, path), sources))
    book = load_workbook(path)
    try:
        assert set(book.sheetnames) == {'summary', 'phase-0', 'phase-1', 'phase-2'}
        assert sorted(book['summary'].values) == sorted([
            ('case', 'value'), ('run-0', 0), ('run-1', 1), ('run-2', 2)])
        for sheet in book:
            assert sheet.freeze_panes == 'B2'
            assert sheet['A1'].font.bold and sheet['A1'].fill.fgColor.rgb
            assert sheet.column_dimensions['A'].width >= 18
            assert sheet.row_dimensions[1].height == 36
            assert sheet.sheet_view.showGridLines is False
    finally:
        book.close()
    before = path.read_bytes()
    write_xlsx(sources[0], [('new phase', ['case'], lambda: iter([('new',)])),
                            ('summary', ['wrong'], lambda: iter([('wrong',)]))])
    with pytest.raises(ValueError, match='header mismatch'):
        publish_workbook(sources[0], path)
    assert path.read_bytes() == before
    assert not list(tmp_path.glob('tmp*.xlsx'))


def test_capacity_table_and_summary_keep_unknown_pools_and_exclude_block_slots():
    from perf_benchmark.app.service_capacity import pool
    from perf_benchmark.plugins.exporter.tabular.render import tables
    from tests.exporter.test_console import Reader

    capacity = dict(engine="vllm", observed_at="2026-10-16T12:00:00+08:00", pools=[
        pool("gpu", "dp:0", 1000, "registration"), pool("gpu", "dp:1", 900, "registration"),
        pool("host_cpu", "shared", None, "LMCache", reason="missing_layout"),
        pool("gpu_block_slots", "engine:0", 99999, "metrics")])
    class CapacityReader(Reader):
        def measurements(self, run_id, *, scope, scope_id):
            return {"evidence": {"service_capacity": capacity}} if scope == "run" else {}

        def request_page(self, run_id, **kwargs):
            return {"items": (), "total": 0}

        def metric_points(self, *args, **kwargs):
            return {}

    specs = tables("run", CapacityReader())
    row = dict(zip(specs[0][1], next(specs[0][2]())))
    assert row["gpu_kv_capacity_tokens (token)"] == 1900
    assert row["host_kv_capacity_tokens (token)"] is None
    assert specs[-1][0] == "kv_capacity"
    assert len(list(specs[-1][2]())) == 4


def test_summary_preserves_each_trials_frozen_price_basis_and_missing_source():
    from tests.exporter.test_console import Reader
    from perf_benchmark.plugins.exporter.tabular.render import tables, compact_summary, units_in_headers

    prices = {"first": {"input_price": 0, "cached_input_price": 0, "output_price": 0,
                        "per_tokens": 1000, "currency": "CNY"},
              "second": {"input_price": 0.5, "cached_input_price": 0.1, "output_price": 2,
                         "per_tokens": 1000000, "currency": "USD", "source": "frozen price list"},
              "historical": {}}
    class Results(Reader):
        def trials(self, run_id):
            return tuple({"trial_id": name, "name": name, "status": "succeeded"} for name in prices)

        def measurements(self, run_id, *, scope, scope_id):
            return {"diagnostics": {"pricing": prices[scope_id]}} if scope == "trial" else {}

        def request_page(self, *args, **kwargs):
            return {"items": (), "total": 0}

        def metric_points(self, *args, **kwargs):
            return {}

    spec = units_in_headers(compact_summary(tables("run", Results())[0]))
    rows = [dict(zip(spec[1], row)) for row in spec[2]()]
    assert rows[0]["input_price"] == 0 and rows[0]["input_price_unit"] == "CNY/1000 tokens"
    assert rows[1]["input_price"] == 0.5 and rows[1]["input_price_unit"] == "USD/1000000 tokens"
    assert [row["per_tokens (tokens)"] for row in rows] == [1000, 1000000, None]
    assert [row["currency"] for row in rows] == ["CNY", "USD", None]
    assert [row["pricing_source"] for row in rows] == [None, "frozen price list", None]
    assert all(rows[2][name] is None for name in ("input_price", "cached_input_price", "output_price"))
    assert isinstance(rows[0]["input_price"], int) and isinstance(rows[1]["input_price"], float)


def test_summary_keeps_maxtpm_latency_cohort_distinct_and_sample_count_units():
    from tests.exporter.test_console import Reader,_point
    from perf_benchmark.plugins.exporter.tabular.render import tables,compact_summary,units_in_headers

    class Results(Reader):
        def request_page(self,*args,**kwargs):
            return {'items':(), 'total':0}

        def metric_points(self,run_id,*,window,**kwargs):
            if window=='maxtpm':
                return {'ttft':_point('s',{'am':2.5,'p50':2,'p90':4,'p99':4,'n':4,'eligible_n':5}),
                        'workload_completion':_point('tok/s',{'value':123})}
            return super().metric_points(run_id,window=window,**kwargs)

    spec=units_in_headers(compact_summary(tables('run',Results())[0]))
    row=dict(zip(spec[1],next(spec[2]())))
    assert row['ttft_mean (s)']==0.11
    assert row['maxtpm_ttft_mean (s)']==2.5
    assert row['maxtpm_ttft_p99 (s)']==4
    assert row['maxtpm_ttft_n (samples)']==4
    assert row['maxtpm_ttft_eligible_n (samples)']==5
    assert row['maxtpm_workload_completion (tok/s)']==123


def test_header_unit_folding_preserves_mixed_units_values_and_missing_reasons():
    from perf_benchmark.plugins.exporter.tabular.render import units_in_headers, PartialValue
    header = ('ttft_mean', 'ttft_p50', 'ttft_unit', 'ttft_availability', 'ttft_reason',
              'cost', 'cost_unit', 'unused_unit')
    rows = [(PartialValue(0.25), None, 's', 'partial', 'missing_timing', 4, 'USD', 's'),
            (0.5, 0.5, 's', 'available', None, 30, 'CNY', 's')]
    _, headings, factory = units_in_headers(('summary', header, lambda: iter(rows)))
    assert headings == ('ttft_mean (s)', 'ttft_p50 (s)', 'ttft_availability', 'ttft_reason',
                        'cost', 'cost_unit')
    assert list(factory()) == [(rows[0][0], None, 'partial', 'missing_timing', 4, 'USD'),
                               (0.5, 0.5, 'available', None, 30, 'CNY')]
    assert isinstance(next(factory())[0], PartialValue)
    long = ('metrics', ('metric', 'value', 'unit'), lambda: iter([('ttft', 0.25, 's')]))
    assert units_in_headers(long)[1] == long[1]
    scalars = ('summary', ('maxtpm', 'maxtpm_unit', 'maxtpm_minutes', 'maxtpm_minutes_unit'),
               lambda: iter([(100, 'tok/min', 3, 'minutes')]))
    assert units_in_headers(scalars)[1] == ('maxtpm (tok/min)', 'maxtpm_minutes (minutes)')


async def test_summary_uses_global_distributions_and_places_parameters_first(tmp_path, stage_context):
    from tests.exporter.test_console import Reader, _point
    from perf_benchmark.plugins.exporter.tabular.render import OVERALL_DISTRIBUTIONS

    class Global(Reader):
        location = str(tmp_path / "run.db")

        def measurements(self, run_id, *, scope, scope_id):
            if scope == "run":
                return {"evidence": {'case':'GLM5.3 cache performance',"conditions": {"model": {"value": "configured-model"},
                        "generation": {"value": {"temperature": 0.7, "api_key": "SECRET"}}}}}
            return {}

        def trials(self, run_id):
            row = dict(super().trials(run_id)[0])
            row["load_profile"] = {"target_concurrency": 8, "max_active_sessions": 16}
            row["parameters"] = {"request_parameters": {"top_p": 0.9, "private": "SECRET"}}
            return (row,)

        def request_page(self, run_id, **kwargs):
            if kwargs["scope"] == "attempt":
                return {"items": (), "total": 0}
            return {"items": ({"measurements": {"request": {"profile": {
                "model": "served-model", "temperature": 0.2, "stream": True}}}},), "total": 1}

        def metric_points(self, run_id, *, scope, scope_id, window, **kwargs):
            if window == "overall":
                return {name: _point("test-unit", {"am": 10, "time_wm": 7, "p50": 20, "p90": 30, "p99": 40})
                        for name in OVERALL_DISTRIBUTIONS if name != "itl"}
            return {"maxtpm": _point("tok/min", {"value": 9000})}

    output = tmp_path / "tables"
    await TabularExporter(parse_settings({"output": str(output)})).export("run", Global(), stage_context("exporter"))
    with ZipFile(output / "results.xlsx") as workbook:
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheet = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
        assert sheet.find("x:sheetData/x:row", ns).get("r") == "1"
        assert sheet.find("x:sheetData/x:row/x:c", ns).get("r") == "A1"
        original = xlsx_rows(workbook, 1)[0]
        assert 'ttft_mean (test-unit)' in original
        assert not any(key.endswith('_unit') for key in original)
        row = without_header_units([original])[0]
        assert list(row)[:7] == ['case',"model", "input_tokens_mean", "output_tokens_mean", "concurrency",
                                "active_users", "generation_parameters"]
        assert row['case']=='GLM5.3 cache performance'
        assert row["model"] == "served-model" and row["concurrency"] == "8" and row["active_users"] == "16"
        assert row["generation_parameters"] == '{"temperature": 0.2, "top_p": 0.9}'
        for name in OVERALL_DISTRIBUTIONS:
            if name == "itl":
                assert "itl_mean" not in row and "itl_p99" not in row
                assert row["itl_reason"] == "metric_missing"
                assert "itl_availability" not in row
            else:
                speed = name in {"prefill_throughput", "decode_throughput"}
                mean = "time_weighted_mean" if speed else "mean"
                assert [row[name + "_" + stat] for stat in (mean, "p50", "p90", "p99")] == ["7" if speed else "10", "20", "30", "40"]
                assert name + "_reason" not in row
        assert float(row["maxtpm"]) == pytest.approx(.009)
        assert 'maxtpm (百万token/M)' in original
        assert row["input_revenue_per_minute"] == ""
        assert row["input_revenue_per_minute_reason"] == "metric_not_recorded"
        assert "maxtpm_start_ns" not in row
        assert all(b"SECRET" not in workbook.read(name) for name in workbook.namelist())
        assert len(workbook.read("xl/workbook.xml").split(b"<sheet ")) - 1 == 3
        assert b'name="metrics"' not in workbook.read("xl/workbook.xml")


@pytest.mark.parametrize("pricing_source", [None, "https://example.org/prices#served-model"])
async def test_sqlite_csv_excel_export_same_minute_revenue_and_request_usage(
    tmp_path, stage_context, pricing_source
):
    pricing = {**PRICES, **({"source": pricing_source} if pricing_source is not None else {})}
    backend = StandardMetrics({"maxtpm": True, "pricing": pricing}).create_backend(seed=0)
    backend.open()
    try:
        backend.consume(
            (
                session(0, 9999),
                *occupancy((1, 1)),
                turn("start", 10, prompt=0, completion=0, cached=0),
                turn("end", 140),
                turn("known", 20, prompt=100, completion=10, cached=25),
                turn("unknown", 150, prompt=100, completion=10),
            )
        )
        report = backend.finalize_trial(trial(4))
    finally:
        backend.close()
    run = Run(id="run", name="test", status="succeeded")
    tr = Trial(id="trial", run_id=run.id, name="=unsafe", status="succeeded")
    tr.measurements.metric_windows = report.windows
    tr.measurements.diagnostics.update(report.diagnostics)
    sess = Session(id="session", trial_id=tr.id, turns=[])
    db = tmp_path / "run.db"
    sink = await SaverRuntime.create(
        SQLiteSaver({"path": str(db)}), run, stage_context("saver", run=run)
    )
    stamp = 10**18 + 1
    try:
        await sink.write_trial(tr)
        await sink.write_session(sess)
        turns = []
        attempts = []
        for index, cache in enumerate((25, None)):
            usage = {"prompt_tokens": 100, "completion_tokens": 10}
            if cache is not None:
                usage["cached_tokens"] = cache
            request = RequestMeasurements(response_completed_ns=stamp + index)
            turns.append(
                frozen_record(TurnResultRecord,
                    id=f"turn-{index}",
                    session_id=sess.id,
                    ordinal=index,
                    status="succeeded",
                    started_at=None,
                    finished_at=None,
                    error=None,
                    estimated_input_tokens=999,
                    reference_output_tokens=None,
                    requested_output_tokens=100,
                    usage=usage,
                    measurements=TurnMeasurements(request=request),
                    metrics={},
                    metadata={"dataset_label": "PRIVATE", "custom_tag": "SECRET"},
                )
            )
            attempts.append(
                frozen_record(AttemptResultRecord,
                    turn_id=f"turn-{index}",
                    session_id=sess.id,
                    attempt=1,
                    status="succeeded",
                    error=None,
                    error_origin=None,
                    measurements=request,
                    usage=usage,
                    metrics={},
                    scenario_passed=None,
                    trigger_ns=None,
                    decision_ns=None,
                    applied_ns=None,
                    plan={},
                )
            )
        sink.write_results(ComputedResultBatch(turns=tuple(turns), attempts=tuple(attempts)))
        await sink.write_metric_points(report.points)
        reader = await sink.finalize(run)
    finally:
        await sink.close()
    output = tmp_path / "tables"
    result = await TabularExporter(parse_settings({"output": str(output), "formats": ["csv", "xlsx"]})).export(
        run.id, reader, stage_context("exporter")
    )
    assert len(result.outputs) == 4
    # The coordinator registers each output by kind alongside HTML's "report".
    artifacts = {"report": str(tmp_path / "report.html")}
    for exported in result.outputs:
        assert exported.kind not in artifacts
        artifacts[exported.kind] = exported.location
    assert set(artifacts) == {
        "report", "tabular:summary.csv", "tabular:minutes.csv",
        "tabular:requests.csv", "tabular:results.xlsx",
    }
    summary = read_csv(output / "summary.csv")[0]
    assert summary["name"] == "'=unsafe"
    assert summary["currency"] == "CNY"
    for name in ("input_price", "cached_input_price", "output_price", "per_tokens"):
        assert float(summary[name]) == PRICES[name]
    assert summary["input_price_unit"] == "CNY/1000 tokens"
    assert summary["per_tokens_unit"] == "tokens"
    assert summary.get("pricing_source") == pricing_source
    assert float(summary["maxtpm"]) == pytest.approx(.000055)
    assert summary["maxtpm_unit"] == '百万token/M'
    native = reader.metric_points(run.id, scope="trial", scope_id=tr.id, window="maxtpm")
    assert native["maxtpm"]["formula_version"] == "dataset-turn-maxtpm-v6-40pct-underfill-60s-first-pass"
    assert native["decode_throughput"]["statistics"]["time_wm"] == 9
    assert float(summary["maxtpm_decode_throughput_time_weighted_mean"]) == 9
    assert float(summary["revenue_per_minute"]) == pytest.approx(0.1175)
    assert summary["revenue_per_minute_unit"] == "CNY/min"
    assert "generation_parameters" not in summary
    assert not {"failed", "cancelled", "success_rate", "stream", "trial_id"} & summary.keys()
    assert summary["total_cached_tokens"] == "25.0*"
    assert "total_cached_tokens_availability" not in summary
    assert 'maxtpm_start_ns' not in summary
    components = {
        'cached_input_revenue_per_minute': 0.0025,
        'input_revenue_per_minute': 0.075,
        'output_revenue_per_minute': 0.04,
    }
    for name, value in components.items():
        assert float(summary[name]) == pytest.approx(value)
        assert summary[name + '_unit'] == 'CNY/min'
    assert sum(float(summary[name]) for name in components) == pytest.approx(float(summary['revenue_per_minute']))
    minutes = read_csv(output / "minutes.csv")
    assert [float(m["revenue_per_minute"]) for m in minutes] == [0.235, 0]
    rows = read_csv(output / "requests.csv")
    assert len(rows) == 4
    for row in rows:
        assert row["response_completed_ns"] == str(stamp + int(row["request_id"].rsplit("-", 1)[1]))
        if row["request_id"] == "turn-0":
            assert row["cached_tokens"] == "25" and row["uncached_input_tokens"] == "75"
            assert "cache_availability" not in row
        else:
            assert row["cached_tokens"] == "" and row["cache_hit_rate"] == ""
            assert row["cache_reason"] == "cache_usage_missing_or_invalid"
    assert not (output / "metrics.csv").exists()
    native_metrics = reader.metric_points(run.id, scope="trial", scope_id=tr.id, window="maxtpm")
    assert native_metrics['revenue_per_minute']['statistics']['value'] == pytest.approx(0.1175)
    with ZipFile(output / "results.xlsx") as z:
        assert b'name="metrics"' not in z.read("xl/workbook.xml")
        for name in z.namelist():
            ET.fromstring(z.read(name))
            assert b"PRIVATE" not in z.read(name) and b"SECRET" not in z.read(name)
        xsummary = without_header_units(xlsx_rows(z, 1))[0]
        assert list(xsummary) == list(without_unit_columns([summary])[0])
        assert xsummary["name"] == "=unsafe"  # Inline text, never an Excel formula.
        assert float(xsummary["revenue_per_minute"]) == float(summary["revenue_per_minute"])
        for name in components:
            assert float(xsummary[name]) == float(summary[name])
        assert xsummary["total_cached_tokens"] == "25.0"
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        cells = sheet.findall("x:sheetData/x:row", ns)[1]
        assert cells[list(xsummary).index("total_cached_tokens")].get("s") == "1"
        for name in ("input_price", "cached_input_price", "output_price", "per_tokens"):
            assert cells[list(xsummary).index(name)].get("t") != "inlineStr"
        assert "input_price (CNY/1000 tokens)" in xlsx_rows(z, 1)[0]
        assert "per_tokens (tokens)" in xlsx_rows(z, 1)[0]
        styles = ET.fromstring(z.read("xl/styles.xml"))
        assert styles.find("x:numFmts/x:numFmt", ns).get("formatCode") == 'General"*"'
        assert without_header_units(xlsx_rows(z, 2)) == without_unit_columns(minutes)
        assert without_header_units(xlsx_rows(z, 3)) == rows
    for file in output.glob("*.csv"):
        assert "PRIVATE" not in file.read_text() and "SECRET" not in file.read_text()
        with file.open(encoding="utf-8-sig", newline="") as stream:
            assert all(key != 'availability' and not key.endswith('_availability')
                       for key in next(csv.reader(stream)))
    # Existing SQLite can also be exported through the offline application entrypoint.
    offline = await export_tables(db, tmp_path / "offline", ["csv"])
    assert len(offline.outputs) == 3
    assert read_csv(tmp_path / "offline" / "summary.csv") == [summary]
    default = await TabularExporter(parse_settings({})).export(run.id, reader, stage_context("exporter"))
    assert [item.location for item in default.outputs] == [str(tmp_path / "run-tables" / "results.xlsx")]
    assert not list((tmp_path / "run-tables").glob("*.csv"))


async def test_partial_cost_is_numeric_in_excel_and_marked_in_csv(tmp_path, stage_context):
    from tests.exporter.test_console import Reader, _point

    class Partial(Reader):
        location = str(tmp_path / "run.db")

        def request_page(self, *args, **kwargs):
            return {"items": (), "total": 0}

        def metric_points(self, *args, **kwargs):
            return {
                "total_cost": _point("CNY", {"value": 0}, availability="partial",
                                     reason="attempt_usage_incomplete"),
                "total_cached_tokens": _point("tokens", {"value": 25}, availability="partial",
                                              reason="partial_usage"),
                "total_completion_tokens": _point("tokens", {"value": 10}),
            }

    output = tmp_path / "tables"
    await TabularExporter(parse_settings({"output": str(output), "formats": ["csv", "xlsx"]})).export(
        "run", Partial(), stage_context("exporter"))
    summary = read_csv(output / "summary.csv")[0]
    assert summary["total_cost"] == "0.0*"
    assert summary["total_cached_tokens"] == "25.0*"
    assert summary["total_completion_tokens"] == "10"
    with ZipFile(output / "results.xlsx") as workbook:
        row = without_header_units(xlsx_rows(workbook, 1))[0]
        assert row["total_cost"] == "0.0" and row["total_cached_tokens"] == "25.0"
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        cells = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml")).findall("x:sheetData/x:row", ns)[1]
        assert cells[list(row).index("total_cost")].get("s") == "1"
        assert cells[list(row).index("total_cost")].get("t") is None
        assert cells[list(row).index("total_completion_tokens")].get("s") is None
    assert not (output / "metrics.csv").exists()


@pytest.mark.parametrize(
    "settings",
    [
        {"formats": []},
        {"formats": ["csv", "csv"]},
        {"formats": ["pdf"]},
        {"output": ""},
        {"unknown": True},
    ],
)
def test_tabular_settings_validation(settings):
    with pytest.raises(ValueError):
        parse_settings(settings)


async def test_tabular_pages_requests_and_preserves_existing_on_read_failure(
    tmp_path, stage_context
):
    from tests.exporter.test_console import Reader

    class Paged(Reader):
        location = str(tmp_path / "run.db")

        def measurements(self, *args, **kwargs):
            return {}

        def request_page(self, run_id, *, scope, offset, limit, **kwargs):
            assert limit == 500
            if scope == "attempt":
                return {"items": (), "total": 0}
            return {
                "items": tuple(
                    {"turn_id": str(i), "usage": {}, "measurements": {}}
                    for i in range(offset, min(offset + limit, 1001))
                ),
                "total": 1001,
            }

    output = tmp_path / "output"
    await TabularExporter(parse_settings({"output": str(output), "formats": ["csv"]})).export(
        "run", Paged(), stage_context("exporter")
    )
    assert len(read_csv(output / "requests.csv")) == 1001
    before = (output / "requests.csv").read_bytes()

    class Failed(Paged):
        def request_page(self, *args, **kwargs):
            if kwargs["offset"]:
                raise RuntimeError("read failed")
            return super().request_page(*args, **kwargs)

    with pytest.raises(RuntimeError, match="read failed"):
        await TabularExporter(parse_settings({"output": str(output)})).export(
            "run", Failed(), stage_context("exporter")
        )
    assert (output / "requests.csv").read_bytes() == before


async def test_summary_keeps_columns_observed_in_any_trial_and_unavailable_reasons(
    tmp_path, stage_context
):
    from tests.exporter.test_console import Reader, _point

    class Mixed(Reader):
        location = str(tmp_path / "run.db")

        def trials(self, run_id):
            template = {**super().trials(run_id)[0], "load_profile": {}}
            return (
                {**template, "trial_id": "known", "parameters": {
                    "request_parameters": {"temperature": 0}}},
                {**template, "trial_id": "unknown"},
            )

        def measurements(self, *args, **kwargs):
            return {"evidence": {"conditions": {"stream": {"value": False}}}}

        def request_page(self, *args, **kwargs):
            return {"items": (), "total": 0}

        def metric_windows(self, *args, **kwargs):
            return ()

        def metric_points(self, run_id, *, scope_id, window, **kwargs):
            if window != "overall":
                return {}
            if scope_id == "known":
                return {"total_cached_tokens": _point("tokens", {"value": 0})}
            return {"total_cached_tokens": {
                "statistics": {"value": None}, "unit": "tokens",
                "availability": "unavailable", "unavailable_reason": "partial_usage"}}

    output = tmp_path / "mixed"
    await TabularExporter(parse_settings({"output": str(output), "formats": ["csv", "xlsx"]})).export(
        "run", Mixed(), stage_context("exporter")
    )
    rows = read_csv(output / "summary.csv")
    assert rows[0]["generation_parameters"] == '{"temperature": 0}'
    assert rows[1]["generation_parameters"] == ""
    assert all("stream" not in row for row in rows)
    assert rows[0]["total_cached_tokens"] == "0"
    assert rows[1]["total_cached_tokens"] == ""
    assert rows[0]["total_cached_tokens_reason"] == ""
    assert rows[1]["total_cached_tokens_reason"] == "partial_usage"
    assert "total_cost" not in rows[0]
    assert "total_cost_availability" not in rows[1]
    assert "concurrency" not in rows[0] and "active_users" not in rows[0]
    with ZipFile(output / "results.xlsx") as workbook:
        assert without_header_units(xlsx_rows(workbook, 1)) == without_unit_columns(rows)
