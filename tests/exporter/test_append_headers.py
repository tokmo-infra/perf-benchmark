"""Workbook appends across data-dependent native report schemas."""
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import load_workbook

from perf_benchmark.plugins.exporter.tabular.append import publish_workbook
from perf_benchmark.plugins.exporter.tabular.render import PartialValue, write_xlsx


def test_append_aligns_reordered_optional_metrics_and_preserves_history(tmp_path):
    path, candidate = tmp_path / 'shared.xlsx', tmp_path / 'candidate.xlsx'
    write_xlsx(path, [('summary', ['case', 'ttft (s)', 'maxtpm_reason'],
                       lambda: iter([('old', 0, 'insufficient_full_minutes')]))])
    write_xlsx(candidate, [('summary', ['case', 'maxtpm (百万token/M)', 'ttft (s)'],
                            lambda: iter([('new', PartialValue(2.5), 0.3)]))])
    source = candidate.read_bytes()
    publish_workbook(candidate, path)
    assert candidate.read_bytes() == source
    # A later run can have only a subset of the accumulated headings.
    write_xlsx(candidate, [('summary', ['case', 'ttft (s)'],
                            lambda: iter([('third', None)]))])
    publish_workbook(candidate, path)
    book = load_workbook(path)
    try:
        sheet = book['summary']
        assert list(sheet.values) == [
            ('case', 'ttft (s)', 'maxtpm_reason', 'maxtpm (百万token/M)'),
            ('old', 0, 'insufficient_full_minutes', None),
            ('new', 0.3, None, 2.5),
            ('third', None, None, None),
        ]
        assert sheet['D3'].number_format == 'General"*"'
        assert sheet['D1'].font.bold and sheet['D1'].fill.fgColor.rgb
        assert sheet.column_dimensions['D'].width >= 18
        assert sheet.freeze_panes == 'B2'
    finally:
        book.close()


def test_append_keeps_different_units_in_distinct_columns(tmp_path):
    path, candidate = tmp_path / 'shared.xlsx', tmp_path / 'candidate.xlsx'
    write_xlsx(path, [('summary', ['case', 'revenue (CNY/min)'], lambda: iter([('a', 2)]))])
    write_xlsx(candidate, [('summary', ['case', 'revenue (USD/min)'], lambda: iter([('b', 3)]))])
    publish_workbook(candidate, path)
    book = load_workbook(path)
    try:
        assert list(book['summary'].values) == [
            ('case', 'revenue (CNY/min)', 'revenue (USD/min)'),
            ('a', 2, None), ('b', None, 3),
        ]
    finally:
        book.close()


@pytest.mark.parametrize('header', [['case', 'value', 'value'], ['case', ''],
                                    ['different', 'value']])
def test_ambiguous_headers_preserve_original_file(tmp_path, header):
    path, candidate = tmp_path / 'shared.xlsx', tmp_path / 'candidate.xlsx'
    write_xlsx(path, [('summary', ['case', 'value'], lambda: iter([('old', 1)]))])
    write_xlsx(candidate, [('summary', header, lambda: iter([['new'] + [2] * (len(header)-1)]))])
    before = path.read_bytes()
    with pytest.raises(ValueError, match='header mismatch'):
        publish_workbook(candidate, path)
    assert path.read_bytes() == before
    assert not list(tmp_path.glob('tmp*.xlsx'))


def test_concurrent_appends_union_optional_columns_under_file_lock(tmp_path):
    path = tmp_path / 'shared.xlsx'
    sources = [tmp_path / f'{i}.xlsx' for i in range(3)]
    for i, source in enumerate(sources):
        write_xlsx(source, [('summary', ['case', f'metric_{i}'],
                             lambda i=i: iter([(f'run-{i}', i)]))])
    with ThreadPoolExecutor(max_workers=3) as pool:
        list(pool.map(lambda source: publish_workbook(source, path), sources))
    book = load_workbook(path)
    try:
        rows = list(book['summary'].values)
        assert len(rows) == 4
        assert set(rows[0]) == {'case', 'metric_0', 'metric_1', 'metric_2'}
        for row in rows[1:]:
            record = dict(zip(rows[0], row))
            index = int(record['case'].split('-')[-1])
            assert record[f'metric_{index}'] == index
            assert all(record[f'metric_{i}'] is None for i in range(3) if i != index)
    finally:
        book.close()


def test_column_union_enforces_excel_limit():
    from perf_benchmark.plugins.exporter.tabular.append import _aligned_columns

    header = ('case', *(f'metric_{i}' for i in range(16_383)))
    assert len(_aligned_columns(header, header, 'summary')) == 16_384
    with pytest.raises(ValueError, match='column limit'):
        _aligned_columns(header, ('case', 'extra_metric'), 'summary')


def test_formula_with_remapped_columns_preserves_original(tmp_path):
    path, candidate = tmp_path / 'shared.xlsx', tmp_path / 'candidate.xlsx'
    write_xlsx(path, [('summary', ['case', 'left', 'right'], lambda: iter([('old', 1, 2)]))])
    write_xlsx(candidate, [('summary', ['case', 'right', 'left'], lambda: iter([('new', 3, 4)]))])
    book = load_workbook(candidate)
    book['summary']['B2'] = '=C2*2'
    book.save(candidate)
    book.close()
    before = path.read_bytes()
    with pytest.raises(ValueError, match='formula requires unchanged column order'):
        publish_workbook(candidate, path)
    assert path.read_bytes() == before


async def test_native_export_appends_missing_and_available_maxtpm(tmp_path, stage_context):
    from perf_benchmark.plugins.exporter.tabular import TabularExporter, parse_settings
    from tests.exporter.test_console import Reader, _point

    class Results(Reader):
        location = str(tmp_path / 'facts.db')

        def request_page(self, run_id, **kwargs):
            return {'items': (), 'total': 0}

        def measurements(self, run_id, *, scope, scope_id):
            return {'evidence': {'case': run_id}} if scope == 'run' else {}

        def metric_points(self, run_id, *, scope, scope_id, window, names=None):
            if window == 'maxtpm':
                if run_id == 'early-stop':
                    return {'maxtpm': _point('tok/min', {'value': None},
                                             availability='unavailable', reason='insufficient_full_minutes')}
                return {'maxtpm': _point('tok/min', {'value': 2_000_000})}
            return super().metric_points(run_id, scope=scope, scope_id=scope_id,
                                         window=window, names=names)

    path = tmp_path / 'comparison.xlsx'
    exporter = TabularExporter(parse_settings({'output': str(tmp_path / 'tables'),
                                               'excel_path': str(path)}))
    reader = Results()
    for name in ('early-stop', 'scored', 'early-stop'):
        await exporter.export(name, reader, stage_context('exporter'))
    book = load_workbook(path)
    try:
        header, *rows = list(book['summary'].values)
        records = [dict(zip(header, row)) for row in rows]
        assert [r['case'] for r in records] == ['early-stop', 'scored', 'early-stop']
        assert [r['maxtpm (百万token/M)'] for r in records] == [None, 2, None]
        assert [r['maxtpm_reason'] for r in records] == [
            'insufficient_full_minutes', None, 'insufficient_full_minutes']
    finally:
        book.close()
