from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from perf_benchmark.plugins.exporter.tabular.append import append_workbook
from perf_benchmark.plugins.exporter.tabular.render import PartialValue, write_xlsx


def test_append_does_not_scan_sheet_dimensions_for_each_row(tmp_path, monkeypatch):
    old, new, joined = (tmp_path / name for name in ('old.xlsx', 'new.xlsx', 'joined.xlsx'))
    write_xlsx(old, [('summary', ['case', 'cost'], lambda: iter([('before', 1)]))])
    rows = [(f'after-{i}', PartialValue(i)) for i in range(200)]
    write_xlsx(new, [('summary', ['case', 'cost'], lambda: iter(rows)),
                     ('new_sheet', ['case', 'cost'], lambda: iter(rows))])
    candidate = load_workbook(new)
    candidate['new_sheet'].conditional_formatting.add(
        'B2:B201', CellIsRule(operator='greaterThan', formula=['100'],
                             fill=PatternFill('solid', fgColor='FCE4D6')))
    candidate.save(new)
    candidate.close()
    getter = Worksheet.max_row.fget
    calls = 0

    def counted(sheet):
        nonlocal calls
        calls += 1
        return getter(sheet)

    with monkeypatch.context() as patch:
        patch.setattr(Worksheet, 'max_row', property(counted))
        append_workbook(old, new, joined)
    assert calls < 30
    book = load_workbook(joined)
    try:
        assert book['summary'].max_row == 202
        assert book['summary']['A3'].value == 'after-0'
        assert book['summary']['B202'].value == 199
        assert book['summary']['B202'].number_format == 'General"*"'
        assert book['new_sheet'].max_row == 201
        assert book['new_sheet']['A1'].value == 'case'
        assert book['new_sheet'].row_dimensions[1].height == 36
        assert book['new_sheet'].freeze_panes == 'B2'
        assert len(book['new_sheet'].conditional_formatting) == 1
    finally:
        book.close()
