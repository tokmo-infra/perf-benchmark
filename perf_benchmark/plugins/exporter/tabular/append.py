"""Publish report sheets atomically, aligning appended metrics by their headers."""
from copy import copy
from pathlib import Path
from tempfile import NamedTemporaryFile
from shutil import copyfile


def _aligned_columns(actual, expected, title):
    """Keep historical columns fixed and append newly observed metric headings."""
    for header in (actual, expected):
        if (not header or any(not isinstance(label, str) or not label.strip()
                              for label in header) or len(set(header)) != len(header)):
            raise ValueError(f"Excel header mismatch in sheet {title}; headers must be unique nonempty text")
    if actual[0] != expected[0]:
        raise ValueError(f"Excel header mismatch in sheet {title}; first column must match")
    combined = dict.fromkeys((*actual, *expected))
    if len(combined) > 16_384:
        raise ValueError(f"Excel column limit exceeded in sheet {title}")
    positions = {label: i for i, label in enumerate(combined, 1)}
    return tuple(positions[label] for label in expected)


def publish_workbook(candidate, destination):
    from filelock import FileLock

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(str(destination) + '.lock'):
        with NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as stream:
            temporary = Path(stream.name)
        try:
            if destination.exists():
                append_workbook(destination, candidate, temporary)
            else:
                copyfile(candidate, temporary)
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)


def append_workbook(existing, candidate, destination):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    old = load_workbook(existing, read_only=False, data_only=False)
    new = None
    try:
        new = load_workbook(candidate, read_only=False, data_only=False)
        mappings = {}
        for sheet in new:
            if sheet.title not in old.sheetnames:
                continue
            target = old[sheet.title]
            rows=sheet.values
            try:
                expected=next(rows,())
            finally:
                rows.close()
            actual = tuple(cell.value for cell in target[1])
            mappings[sheet.title] = _aligned_columns(actual, expected, sheet.title)
            if target.max_row + max(0, sheet.max_row - 1) > 1_048_576:
                raise ValueError(f"Excel row limit exceeded in sheet {sheet.title}")
        for sheet in new:
            present = sheet.title in old.sheetnames
            target = old[sheet.title] if present else old.create_sheet(sheet.title)
            columns = mappings[sheet.title] if present else tuple(range(1, sheet.max_column + 1))
            remapped = columns != tuple(range(1, len(columns) + 1))
            previous_columns = target.max_column if present else 0
            if not present:
                target.sheet_format = copy(sheet.sheet_format)
                target.sheet_properties = copy(sheet.sheet_properties)
                target.views = copy(sheet.views)
                target.auto_filter = copy(sheet.auto_filter)
                target.conditional_formatting = copy(sheet.conditional_formatting)
                for key, dimension in sheet.column_dimensions.items():
                    target.column_dimensions[key] = copy(dimension)
            for source_column, target_column in enumerate(columns if present else (), 1):
                if target_column <= previous_columns:
                    continue
                source = sheet.cell(1, source_column)
                cell = target.cell(1, target_column, source.value)
                cell.font, cell.fill = copy(source.font), copy(source.fill)
                cell.border, cell.alignment = copy(source.border), copy(source.alignment)
                cell.number_format, cell.protection = source.number_format, copy(source.protection)
                key = get_column_letter(source_column)
                if key in sheet.column_dimensions:
                    dimension = copy(sheet.column_dimensions[key])
                    dimension.index = get_column_letter(target_column)
                    dimension.min = dimension.max = target_column
                    target.column_dimensions[dimension.index] = dimension
            # max_row scans all populated cells; never call it per appended row.
            target_index = target.max_row + 1 if present else 1
            for row in sheet.iter_rows(min_row=2 if present else 1):
                values = [None] * max(columns)
                for source, column in zip(row, columns):
                    if source.data_type == 'f' and remapped:
                        raise ValueError(f"Excel formula requires unchanged column order in sheet {sheet.title}")
                    values[column - 1] = source.value
                target.append(values)
                source_index = row[0].row
                if source_index in sheet.row_dimensions:
                    dimension = copy(sheet.row_dimensions[source_index])
                    dimension.index = target_index
                    target.row_dimensions[target_index] = dimension
                for source, column in zip(row, columns):
                    cell = target.cell(target_index, column)
                    if source.has_style:
                        cell.font = copy(source.font)
                        cell.fill = copy(source.fill)
                        cell.border = copy(source.border)
                        cell.alignment = copy(source.alignment)
                        cell.number_format = source.number_format
                        cell.protection = copy(source.protection)
                target_index += 1
            target.freeze_panes = 'B2'
            from perf_benchmark.plugins.exporter.tabular.render import text_width
            target.column_dimensions['A'].width = min(255, max(
                target.column_dimensions['A'].width or 18,
                max((text_width(c.value) + 3 for c in target['A'] if c.value is not None), default=18)))
        old.save(Path(destination))
    finally:
        old.close()
        if new is not None:
            new.close()
